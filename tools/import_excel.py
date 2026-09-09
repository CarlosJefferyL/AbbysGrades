#!/usr/bin/env python3
"""
Convierte el libro de Excel "Calificaciones_Semestre_*.xlsx" (formato que se
usaba antes de esta app) en un respaldo .json que se puede cargar con el
botón "Restaurar" de la aplicación.

Uso:
    pip install openpyxl
    python3 tools/import_excel.py Calificaciones.xlsx salida.json [--expected expected.json]

Formato esperado del libro (por cada nivel, p. ej. "I"):
  - Hoja "Ingles I":            columnas Examen 1..3 (0-1) y Examen Final.
  - Hoja "Asistencia Ingles I": bloque de asistencia (On Time / Tardy / Absent)
                                y bloque de participación (P / X) con las
                                mismas fechas; el rango de columnas de cada
                                parcial se lee de las fórmulas COUNTIF.
  - Hoja "Tareas I":            una columna por tarea (0-100), separadas por
                                las columnas "TC Parcial N".

Con --expected también escribe los valores ya calculados por Excel para
verificar la app (ver tools/compare_with_excel.js).
"""
import json
import re
import sys
import uuid
from datetime import datetime, date

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Falta openpyxl: pip install openpyxl")

WEIGHTS = {"attendance": 15, "participation": 15, "homework": 20, "exam": 50}
ATT_MAP = {"on time": "P", "tardy": "R", "absent": "A"}
PART_MAP = {"p": "P", "x": "X"}


def uid():
    return uuid.uuid4().hex[:8]


def norm_name(s):
    s = re.sub(r"\s*\d+\s*$", "", str(s or ""))  # quita el número al final ("Mya Gonzalez 6")
    return re.sub(r"\s+", " ", s).strip().lower()


def clean_name(s):
    return re.sub(r"\s+", " ", re.sub(r"\s*\d+\s*$", "", str(s or ""))).strip()


def as_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return None


def pct(v):
    """Excel guarda 0.69 -> app usa 69."""
    if isinstance(v, (int, float)):
        return round(float(v) * 100, 4)
    return None


def find_row(ws, col, text, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip().lower().startswith(text.lower()):
            return r
    return None


def parse_ranges(formula_ws, row):
    """Lee los rangos de columnas de cada parcial de las fórmulas COUNTIF(E3:AB3, ...)."""
    ranges = []
    for c in (2, 3, 4):
        f = formula_ws.cell(row, c).value
        m = re.search(r"\(([A-Z]+)\d+:([A-Z]+)\d+", str(f))
        if not m:
            raise ValueError(f"No se pudo leer el rango del parcial en {formula_ws.title}!{get_column_letter(c)}{row}: {f}")
        ranges.append((column_index_from_string(m.group(1)), column_index_from_string(m.group(2))))
    return ranges


def student_rows(ws, header_row, name_col=1):
    """Filas de alumnos debajo de un encabezado 'Nombre' hasta la primera fila vacía."""
    rows = {}
    r = header_row + 1
    while r <= ws.max_row:
        v = ws.cell(r, name_col).value
        if not isinstance(v, str) or not v.strip():
            break
        rows[norm_name(v)] = r
        r += 1
    return rows


def import_level(wf, wv, level, report):
    g = wf[f"Ingles {level}"]
    gv = wv[f"Ingles {level}"]
    a = wf[f"Asistencia Ingles {level}"]
    av = wv[f"Asistencia Ingles {level}"]
    t = wf[f"Tareas {level}"]
    tv = wv[f"Tareas {level}"]

    # --- alumnos (la hoja de calificaciones es la lista maestra)
    students = []
    grade_rows = {}
    for r in range(2, g.max_row + 1):
        name = g.cell(r, 1).value
        if isinstance(name, str) and name.strip():
            sid = uid()
            students.append({"id": sid, "name": clean_name(name), "code": ""})
            grade_rows[norm_name(name)] = (sid, r)

    # --- encabezados de la hoja de calificaciones
    headers = {str(c.value).strip().lower(): c.column for c in g[1] if c.value}
    exam_cols = [headers.get(f"examen {i}") for i in (1, 2, 3)]
    final_exam_col = headers.get("examen final")

    # --- asistencia y participación
    att_header = find_row(a, 1, "Nombre")
    part_header = find_row(a, 1, "Nombre", att_header + 1)
    att_rows = student_rows(a, att_header)
    part_rows = student_rows(a, part_header)
    ranges = parse_ranges(a, att_header + 1)

    partials = []
    for pi, (c1, c2) in enumerate(ranges, start=1):
        p = {
            "id": uid(), "name": f"Parcial {pi}", "sessions": [],
            "attendance": {"marks": {}}, "participation": {"marks": {}},
            "homework": {"items": [], "scores": {}},
            "exam": {"items": [{"id": uid(), "name": "Examen", "max": 100}], "scores": {}},
        }
        # Sesiones: columnas del rango con fecha y al menos una marca reconocida.
        for c in range(c1, c2 + 1):
            d = as_date(av.cell(att_header, c).value)
            if not d:
                continue
            has_mark = any(
                isinstance(a.cell(r, c).value, str) and a.cell(r, c).value.strip().lower() in ATT_MAP
                for r in att_rows.values()
            )
            if has_mark:
                p["sessions"].append({"id": uid(), "date": d, "_col": c})
        for key, (sid, _) in grade_rows.items():
            ar = att_rows.get(key)
            pr = part_rows.get(key)
            if ar is None:
                report.append(f"[{level}] {key}: no está en la hoja de asistencia")
            if pr is None:
                report.append(f"[{level}] {key}: no está en el bloque de participación")
            am, pm = {}, {}
            for s in p["sessions"]:
                c = s["_col"]
                if ar is not None:
                    v = a.cell(ar, c).value
                    k = str(v).strip().lower() if v is not None else ""
                    if k in ATT_MAP:
                        am[s["id"]] = ATT_MAP[k]
                    else:
                        # En el Excel, cualquier otra cosa (vacío, notas) no contaba como asistencia.
                        am[s["id"]] = "A"
                        if k:
                            report.append(f"[{level}] asistencia {key} {s['date']}: valor '{v}' -> Ausente")
                if pr is not None:
                    v = a.cell(pr, c).value
                    k = str(v).strip().lower() if v is not None else ""
                    if k in PART_MAP:
                        pm[s["id"]] = PART_MAP[k]
                    else:
                        pm[s["id"]] = "X"
                        if k:
                            report.append(f"[{level}] participación {key} {s['date']}: valor '{v}' -> X")
            p["attendance"]["marks"][sid] = am
            p["participation"]["marks"][sid] = pm
        for s in p["sessions"]:
            del s["_col"]
        partials.append(p)

    # --- tareas: columnas entre los "TC Parcial N"
    task_rows = student_rows(t, 1)
    pi = 0
    for c in range(2, t.max_column + 1):
        h = t.cell(1, c).value
        if h is None:
            continue
        hs = str(h).strip()
        if re.match(r"^TC\s*Parcial", hs, re.I):
            pi += 1
            continue
        if pi >= len(partials):
            break
        item = {"id": uid(), "name": hs, "max": 100}
        partials[pi]["homework"]["items"].append(item)
        for key, (sid, _) in grade_rows.items():
            tr = task_rows.get(key)
            if tr is None:
                continue
            v = tv.cell(tr, c).value
            if isinstance(v, (int, float)):
                partials[pi]["homework"]["scores"].setdefault(sid, {})[item["id"]] = float(v)
    for key in grade_rows:
        if key not in task_rows:
            report.append(f"[{level}] {key}: no está en la hoja de tareas")

    # --- exámenes y examen final (hoja de calificaciones)
    final_exam = {}
    for key, (sid, r) in grade_rows.items():
        for pi, col in enumerate(exam_cols):
            v = pct(gv.cell(r, col).value) if col else None
            if v is not None:
                item_id = partials[pi]["exam"]["items"][0]["id"]
                partials[pi]["exam"]["scores"].setdefault(sid, {})[item_id] = v
        fe = pct(gv.cell(r, final_exam_col).value) if final_exam_col else None
        if fe is not None:
            final_exam[sid] = fe

    group = {
        "id": uid(), "name": f"Inglés {level}", "weights": dict(WEIGHTS),
        "passing": 70, "finalExamWeight": 50, "tardyValue": 1, "blankAsZero": False,
        "students": students, "partials": partials, "finalExam": final_exam,
    }

    # --- valores calculados por Excel, para verificar la app
    expected = {}
    for key, (sid, r) in grade_rows.items():
        def col(name):
            c = headers.get(name)
            return pct(gv.cell(r, c).value) if c else None
        expected[sid] = {
            "name": clean_name(g.cell(r, 1).value),
            "partials": [
                {
                    "attendance": col(f"asistencia {i}"),
                    "participation": col(f"eval oral {i}"),
                    "homework": col(f"tareas parcial {i}"),
                    "exam": col(f"examen {i}"),
                    "grade": col(f"parcial {i}"),
                }
                for i in (1, 2, 3)
            ],
            "average": col("promedio parciales") if "promedio parciales" in headers else col("prom parciales"),
            "needed": col("calif necesaria"),
        }
    return group, expected


def main(argv):
    if len(argv) < 3:
        sys.exit(__doc__)
    src, dst = argv[1], argv[2]
    expected_path = argv[argv.index("--expected") + 1] if "--expected" in argv else None
    wf = openpyxl.load_workbook(src, data_only=False)
    wv = openpyxl.load_workbook(src, data_only=True)
    levels = [m.group(1) for ws in wf.worksheets for m in [re.match(r"^Ingles (\S+)$", ws.title)] if m]
    report = []
    groups, expected = [], {}
    for level in levels:
        group, exp = import_level(wf, wv, level, report)
        groups.append(group)
        expected[group["id"]] = exp
        print(f"Inglés {level}: {len(group['students'])} alumnos, "
              + ", ".join(f"{p['name']}: {len(p['sessions'])} sesiones / {len(p['homework']['items'])} tareas" for p in group["partials"]))
    data = {"version": 1, "groups": groups, "activeGroupId": groups[0]["id"] if groups else None}
    with open(dst, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    print(f"Escrito {dst}")
    if expected_path:
        with open(expected_path, "w", encoding="utf-8") as fh:
            json.dump(expected, fh, ensure_ascii=False, indent=1)
        print(f"Escrito {expected_path}")
    if report:
        print("\nAvisos:")
        for line in report:
            print("  " + line)


if __name__ == "__main__":
    main(sys.argv)
